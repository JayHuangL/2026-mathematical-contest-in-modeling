"""Fallback workbook exporter when the optional @oai/artifact-tool is unavailable.

The four build_workbook.mjs files remain the preferred exporter.  This small
fallback copies each untouched template from 附件/附件3 and replaces only the
computed value area with the current result*_data.json payload.
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import openpyxl


HERE = Path(__file__).resolve().parent
TEMPLATE_DIR = HERE / "附件" / "附件3"
OUTPUT_DIR = HERE / "outputs"


def write_values(sheet, data, value_key: str, include_surface: bool = False) -> None:
    if include_surface:
        headers = ["时间\\到药材中心的距离", *data["fixed_r_cm"], "药材表面"]
    else:
        headers = ["时间\\到药材中心的距离", *data["r_cm"]]
    last_row = len(data["t_s"]) + 1
    sheet.cell(1, 1).value = headers[0]
    for col, value in enumerate(headers[1:], 2):
        sheet.cell(1, col).value = value
    values = data[value_key]
    for row_index, (time_s, row_values) in enumerate(zip(data["t_s"], values), 2):
        sheet.cell(row_index, 1).value = time_s
        for col, value in enumerate(row_values, 2):
            sheet.cell(row_index, col).value = value

    for row in sheet.iter_rows(min_row=2, max_row=last_row,
                               min_col=1, max_col=len(headers)):
        row[0].number_format = "0.0000" if row[0].row == last_row else "0"
        for cell in row[1:]:
            cell.number_format = "0.0000"
    for cell in sheet[1][:len(headers)]:
        cell.number_format = "0.0" if cell.column > 1 else "General"
    sheet.freeze_panes = "B2"


def export(question: str, template_name: str, output_name: str,
           data_name: str, sheet_keys: tuple[tuple[str, str], ...],
           include_surface: bool = False) -> None:
    data = json.loads((HERE / question / data_name).read_text(encoding="utf-8"))
    output_dir = OUTPUT_DIR / output_name
    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / template_name
    shutil.copy2(TEMPLATE_DIR / template_name, target)
    book = openpyxl.load_workbook(target)
    for sheet_name, value_key in sheet_keys:
        write_values(book[sheet_name], data, value_key, include_surface)
    book.save(target)


def main() -> None:
    export("第一问", "result1.xlsx", "q1", "result1_data.json",
           (("温度", "T"), ("水分浓度", "C")))
    shutil.copy2(OUTPUT_DIR / "q1" / "result1.xlsx", HERE / "第一问" / "result1.xlsx")

    export("第二问", "result2.xlsx", "q2", "result2_data.json",
           (("温度", "T"), ("水分浓度", "C")))

    export("第三问", "result3.xlsx", "q3", "result3_data.json",
           (("Sheet1", "C"),))
    export("第四问", "result4.xlsx", "q4", "result4_data.json",
           (("Sheet1", "C"),), include_surface=True)


if __name__ == "__main__":
    main()
