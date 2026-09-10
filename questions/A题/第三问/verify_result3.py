from pathlib import Path
import argparse,json,hashlib
import numpy as np
import openpyxl
HERE=Path(__file__).resolve().parent
p=argparse.ArgumentParser()
p.add_argument('--root',type=Path,default=HERE.parent)
p.add_argument('--book',type=Path)
a=p.parse_args()
book=a.book or a.root/'outputs/q3/result3.xlsx'
s=json.loads((HERE/'validation.json').read_text(encoding='utf-8'))
payload=json.loads((HERE/'result3_data.json').read_text(encoding='utf-8'))
assert hashlib.sha256((a.root/'附件/附件1.xlsx').read_bytes()).hexdigest()==s['input_sha256']
assert hashlib.sha256((a.root/'第二问/solve_q2.py').read_bytes()).hexdigest()==s['q2_code_sha256']
with np.load(HERE/'result3_full_precision.npz') as full:
    expected=np.round(full['C'][1:],4)
    assert np.array_equal(expected,payload['C'])
    assert np.array_equal(np.round(full['t'][1:],4),payload['t_s'])
    assert np.all(np.diff(full['t'][:-1])==60)
    assert full['C'][-1].max()<.15 and full['C'][-2].max()>.15
    assert full['t'][-2]<s['event_s']<full['t'][-1]
    for t,c in zip(s['table_t_s'],s['table_C']):
        i=int(np.argmin(abs(full['t']-t)))
        assert np.array_equal(full['C'][i,::5],c)
w=openpyxl.load_workbook(book,data_only=True,read_only=True)
assert w.sheetnames==['Sheet1']
rows=iter(w.active.iter_rows())
header=next(rows)
assert [c.value for c in header[1:]]==payload['r_cm']
count=0
for i,row in enumerate(rows):
    assert row[0].value==payload['t_s'][i]
    assert np.array_equal([c.value for c in row[1:]],expected[i])
    assert all(c.number_format=='0.0000' for c in row[1:])
    count+=1
assert count==s['output_time_rows']
w.close()
report=(HERE/'第三问公式与求解.md').read_text(encoding='utf-8')
table=(HERE/'结果表.md').read_text(encoding='utf-8')
assert '{{' not in report
assert all(line in report for line in table.splitlines() if line.startswith('|'))
result={'time_rows':count,'numeric_cells':count*21,'all_values_match':True,'strict_terminal_check':True,'input_and_q2_unchanged':True}
(HERE/'validation_export.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
