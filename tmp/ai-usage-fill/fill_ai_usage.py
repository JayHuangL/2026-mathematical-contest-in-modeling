import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips


REFERENCE = Path(r"D:\Code\2026-mathematical-contest-in-modeling\paper\AI 工具使用详情.docx")
OUTPUT = Path(r"D:\Code\2026-mathematical-contest-in-modeling\paper\AI 工具使用详情（已填写）.docx")


def set_run_font(run, size=9.2, name="宋体", bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)


def clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def set_paragraph_text(paragraph, text, *, align=None, size=10.5, bold=True):
    clear_paragraph(paragraph)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return paragraph


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Twips(width)


def set_table_widths(table, widths):
    table.autofit = False
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for col_index, width in enumerate(widths):
            set_cell_width(row.cells[col_index], width)


def set_cell_text(cell, text, *, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.0, bold=False):
    # Keep the cell properties and the first paragraph's original properties.
    paragraphs = list(cell.paragraphs)
    for paragraph in paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    paragraph = paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    parts = text.split("\n")
    for part_index, part in enumerate(parts):
        if part_index:
            paragraph.runs[-1].add_break()
        run = paragraph.add_run(part)
        set_run_font(run, size=size, bold=bold)


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def fill_detail_table(table, values):
    # values maps row index to the right-hand cell text.
    set_table_widths(table, [2200, 6870])
    for row_index, value in values.items():
        set_cell_text(table.cell(row_index, 1), value, align=WD_ALIGN_PARAGRAPH.LEFT, size=8.8)
    for row_index in range(len(table.rows)):
        set_cell_text(table.cell(row_index, 0), table.cell(row_index, 0).text, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.2)
    set_cell_text(table.cell(0, 0), "项目", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, bold=True)
    set_cell_text(table.cell(0, 1), "填写内容", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, bold=True)


def main():
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)

    # Remove instructions from the retained form after recording the usable slots.
    original_paragraphs = list(doc.paragraphs)
    for idx in (16, 13, 8):
        remove_paragraph(original_paragraphs[idx])

    current_paragraphs = list(doc.paragraphs)
    # The original labels were immediately before the three detail tables.
    set_paragraph_text(current_paragraphs[8], "交互记录1：", size=10.5, bold=True)
    set_paragraph_text(current_paragraphs[9], "交互记录2：", size=10.5, bold=True)
    set_paragraph_text(current_paragraphs[10], "交互记录3：", size=10.5, bold=True)
    current_paragraphs[8].paragraph_format.keep_with_next = True
    current_paragraphs[9].paragraph_format.keep_with_next = True
    current_paragraphs[10].paragraph_format.keep_with_next = True

    # Keep the overview tables grounded in the existing template record, while making the purposes specific to A题.
    set_cell_text(doc.tables[0].cell(1, 3), "赛题理解，辅助热湿模型与移动边界建模，编写和调试数值代码、代码注释，分析算法可行性", size=9.0)
    set_cell_text(doc.tables[0].cell(2, 3), "中文语言润色，参考文献格式整理", size=9.0)
    set_cell_text(doc.tables[1].cell(3, 3), "分析固定半径、变物性和收缩移动边界模型的可行性", size=9.0)
    set_cell_text(doc.tables[1].cell(4, 3), "辅助编写有限体积和隐式BDF求解代码，调试报错", size=9.0)

    fill_detail_table(doc.tables[3], {
        1: "模型建立与算法设计",
        2: "Codex（GPT-5.6 luna）",
        3: "2026-09-11",
        4: "请结合A题、附件1和附录2，分析问题1的建模思路：如何将圆柱形药材简化为径向一维模型？环境温度和水分浓度如何连续拟合？请比较候选拟合方法，并说明有限体积法、隐式BDF及中心和表面边界条件的处理。",
        5: "建议采用固定半径轴对称圆柱的径向热传导与水分扩散方程；先用时间阻塞交叉验证比较分段线性、PCHIP、Akima、样条和拉伸指数模型，综合RMSE与平滑性选择边界函数；空间离散用有限体积法，时间推进用隐式BDF，中心为零通量，表面采用Robin换热和传质边界。",
        6: "部分修改后采纳：采用一维径向热湿模型、拉伸指数边界拟合、有限体积法和隐式BDF；根据题目参数补充单位换算、Kirchhoff界面平均以及网格和容差设置，未直接照搬AI给出的参数。",
        7: "队伍独立推导能量守恒和水分守恒方程，按附录2核对参数与单位；对候选边界拟合进行时间阻塞交叉验证，并检查网格收敛、表面通量和质量守恒。",
    })

    fill_detail_table(doc.tables[4], {
        1: "模型假设与符号定义",
        2: "Codex（GPT-5.6 luna）",
        3: "2026-09-13",
        4: "请检查A题问题1至问题4中模型假设、符号和边界条件是否一致，尤其核对C为干基含水率、经验公式中的T用K、r与R(t)的单位，以及问题4由移动半径映射到固定材料坐标时的守恒关系。请列出需要修改的符号、单位和推导风险。",
        5: "应统一温度的摄氏度与开尔文使用、含水率C的单位和符号；中心采用轴对称零通量，表面采用有限速率Robin条件。问题4可令材料坐标为ξ=r/R(t)，并用材料导数处理移动边界；还应检查半径插值的单调性及外部位置输出。",
        6: "部分修改后采纳：依据题目附录和独立推导统一T、C、r、R(t)、ξ等符号；保留问题4的材料坐标及整体收支关系，重写部分表述和公式，未采纳未经题目参数支持的假设。",
        7: "逐项进行量纲和符号核对；独立推导材料导数、边界通量和平均含水率收支式，并在代码中检查中心和表面边界以及收缩后的有效区域。",
    })

    fill_detail_table(doc.tables[5], {
        1: "辅助整理引用文献格式",
        2: "DeepSeek（V4.1 flash）",
        3: "2026-09-13",
        4: "请将论文末尾的7条参考文献按GB/T 7714格式统一，保留我提供的作者、题名、文献类型、期刊或会议、年份、卷期、页码和DOI；不补写缺失信息，并检查正文引用编号与文后条目是否对应。",
        5: "建议按统一格式整理期刊文献“作者.题名[J].期刊, 年, 卷(期):页码.DOI”，会议论文使用[C]//并保留DOI；整理后仍需逐条核对作者、出版信息和文献类型。",
        6: "部分修改后采纳：统一标点、作者格式及[J]/[C]标识，队伍逐条核对作者、题名、出版信息和DOI并修正；未使用AI补充缺失的文献信息。",
        7: "将7条文献与原始条目和正文引用逐项比对，检查文献类型、年份卷期、页码、DOI及编号对应关系。",
    })

    # Overall adoption and verification record.
    set_table_widths(doc.tables[6], [600, 1700, 3385, 3385])
    overall = [
        ("部分修改后采纳：参考AI提出的热传导-水分扩散耦合框架，结合题目边界和附录参数，建立一维径向有限体积模型；问题四由队伍进一步引入材料坐标和移动边界。", "队伍独立推导控制方程和初边值条件，并用网格收敛、质量收支和敏感性分析检验。"),
        ("部分修改后采纳：AI用于检查守恒方程、Robin边界、Kirchhoff界面平均和材料导数推导；公式按题目附录与队伍推导重写，未直接照搬。", "逐项核对量纲和符号，独立推导关键公式，并用代码复现界面通量和Jacobian。"),
        ("部分修改后采纳：AI辅助组织有限体积、隐式BDF、稀疏Jacobian、事件检测和结果导出代码；队伍按实际网格、步长和文件格式修改。", "独立运行并复核初边值条件、正值性、径向单调性、网格与时间收敛、守恒残差及result1至result4输出。"),
        ("未采纳AI直接生成的结果分析或模型评价，相关内容由队伍依据计算结果和独立检验完成。", "重新读取结果文件并复算关键时刻，核对57.6790 h、51.2683 h、0.1500 kg/kg等数值与表图及误差分析。"),
        ("部分修改后采纳（仅文字辅助）：AI协助润色摘要、问题分析和结论，模型结构、数值结果及结论由队伍提供并改写。", "逐句对照题目要求、计算表和图，队员交叉审阅并删除未经结果支持的表述。"),
        ("部分修改后采纳：DeepSeek辅助统一7条参考文献的格式；作者、题名、期刊或会议、年份、卷期、页码和DOI由队伍逐条保留并核对。", "检查文献字段完整性、标点和编号，并与正文引用逐项对应。"),
    ]
    for row_index, (adoption, verification) in enumerate(overall, start=1):
        set_cell_text(doc.tables[6].cell(row_index, 2), adoption, size=8.4)
        set_cell_text(doc.tables[6].cell(row_index, 3), verification, size=8.4)
    for row_index in range(len(doc.tables[6].rows)):
        set_cell_text(doc.tables[6].cell(row_index, 0), doc.tables[6].cell(row_index, 0).text, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.2, bold=(row_index == 0))
        set_cell_text(doc.tables[6].cell(row_index, 1), doc.tables[6].cell(row_index, 1).text, align=WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT, size=9.2, bold=(row_index == 0))
    set_cell_text(doc.tables[6].cell(0, 2), "采纳与修改情况", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, bold=True)
    set_cell_text(doc.tables[6].cell(0, 3), "人工核验方式", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, bold=True)

    set_table_widths(doc.tables[7], [600, 2100, 1200, 5170])
    contributions = [
        "是",
        "是",
        "是",
        "是",
        "是",
    ]
    descriptions = [
        "队伍根据题目建立一维圆柱径向热湿模型；问题二引入变物性耦合，问题四引入收缩移动边界和材料坐标，并确定边界拟合方案。",
        "队伍独立推导能量和水分守恒、Robin边界、有限体积离散、Kirchhoff界面系数及移动边界收支关系，确定BDF和Jacobian处理。",
        "队伍确定网格、步长、容差、稳定判据、事件检测、收敛和守恒检查及结果导出；AI仅辅助代码实现和调试。",
        "队伍独立读取结果并复核关键数值，完成收敛、守恒和敏感性分析，撰写模型评价与结论。",
        "队伍完成表格、结果文件、图表制作和论文定稿；AI仅用于文字润色和参考文献格式辅助。",
    ]
    for row_index in range(len(doc.tables[7].rows)):
        for col_index in (0, 1, 2):
            set_cell_text(doc.tables[7].cell(row_index, col_index), doc.tables[7].cell(row_index, col_index).text, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.2, bold=(row_index == 0))
    for row_index, (yes_no, description) in enumerate(zip(contributions, descriptions), start=1):
        set_cell_text(doc.tables[7].cell(row_index, 2), yes_no, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.2)
        set_cell_text(doc.tables[7].cell(row_index, 3), description, align=WD_ALIGN_PARAGRAPH.LEFT, size=8.8)
    set_cell_text(doc.tables[7].cell(0, 3), "本队贡献说明", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5, bold=True)

    # Keep the section title and confirmation sentence visually attached to their tables.
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(("四、", "核心环节人工主导确认")):
            paragraph.paragraph_format.keep_with_next = True

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
