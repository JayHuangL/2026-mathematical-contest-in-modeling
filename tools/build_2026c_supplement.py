from pathlib import Path
import shutil

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\Code\2026-mathematical-contest-in-modeling")
SOURCE = ROOT / "paper" / "2026C论文.docx"
OUTPUT = ROOT / "paper" / "2026C论文_补充完善版.docx"
FLOW_DIR = ROOT / "paper" / "tmp" / "20260913_flowcharts"
FLOW_DIR.mkdir(parents=True, exist_ok=True)


def clear_paragraph(paragraph):
    """Remove paragraph content while retaining its paragraph properties."""
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_paragraph(paragraph, text):
    clear_paragraph(paragraph)
    paragraph.add_run(text)
    return paragraph


def insert_paragraph_after(paragraph, text=None, style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style and style in [s.name for s in paragraph.part.document.styles]:
        new_paragraph.style = style
    if text is not None:
        new_paragraph.add_run(text)
    return new_paragraph


def delete_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)
    paragraph._p = paragraph._element = None


def find_paragraph(doc, text, exact=False):
    for paragraph in doc.paragraphs:
        value = paragraph.text.strip()
        if (value == text) if exact else (text in value):
            return paragraph
    raise ValueError(f"找不到段落：{text}")


def paragraph_index(doc, target):
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph._p is target._p:
            return i
    raise ValueError("目标段落不在文档中")


def set_cell(cell, value, size=9):
    cell.text = str(value)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)


def add_flowchart(doc, anchor, image_path, caption):
    image_paragraph = insert_paragraph_after(anchor)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.15))
    caption_paragraph = insert_paragraph_after(image_paragraph)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.style = "图注" if "图注" in [s.name for s in doc.styles] else "Normal"
    run = caption_paragraph.add_run(caption)
    run.font.size = Pt(9)
    return caption_paragraph


def choose_font(size):
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
        Path(r"C:\Windows\Fonts\Deng.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                return ImageFont.truetype(str(candidate), size, index=0)
            except Exception:
                try:
                    return ImageFont.truetype(str(candidate), size)
                except Exception:
                    pass
    return ImageFont.load_default()


def wrap_text(draw, text, font, max_width):
    lines = []
    for part in text.split("\n"):
        current = ""
        for char in part:
            candidate = current + char
            box = draw.textbbox((0, 0), candidate, font=font)
            if current and box[2] - box[0] > max_width:
                lines.append(current)
                current = char
            else:
                current = candidate
        if current:
            lines.append(current)
    return lines


def draw_flowchart(path, title, nodes):
    width, height = 2200, 600
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = choose_font(38)
    node_font = choose_font(28)
    title_box = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((width - (title_box[2] - title_box[0])) / 2, 32), title, fill=(35, 55, 75), font=title_font)

    node_width = 270
    node_height = 250
    margin = 40
    gap = (width - 2 * margin - len(nodes) * node_width) / (len(nodes) - 1)
    top = 145
    fills = [(237, 245, 253), (247, 243, 231), (240, 246, 239)]
    outline = (57, 91, 123)

    positions = []
    for i, node in enumerate(nodes):
        x = int(margin + i * (node_width + gap))
        y = top
        positions.append((x, y))
        fill = fills[i % len(fills)]
        draw.rounded_rectangle((x, y, x + node_width, y + node_height), radius=18, fill=fill, outline=outline, width=3)
        lines = wrap_text(draw, node, node_font, node_width - 34)
        line_heights = []
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=node_font)
            line_heights.append(bbox[3] - bbox[1])
        total_height = sum(line_heights) + 10 * max(0, len(lines) - 1)
        current_y = y + (node_height - total_height) / 2
        for line, line_height in zip(lines, line_heights):
            bbox = draw.textbbox((0, 0), line, font=node_font)
            text_width = bbox[2] - bbox[0]
            draw.text((x + (node_width - text_width) / 2, current_y), line, fill=(35, 55, 75), font=node_font)
            current_y += line_height + 10

    arrow_y = top + node_height / 2
    for i in range(len(positions) - 1):
        x1 = positions[i][0] + node_width + 5
        x2 = positions[i + 1][0] - 8
        draw.line((x1, arrow_y, x2, arrow_y), fill=(82, 112, 139), width=5)
        draw.polygon([(x2, arrow_y), (x2 - 17, arrow_y - 11), (x2 - 17, arrow_y + 11)], fill=(82, 112, 139))

    image.save(path, dpi=(200, 200))


def fill_section(doc, heading_text, next_heading_text, items):
    heading = find_paragraph(doc, heading_text, exact=True)
    paragraphs = doc.paragraphs
    start = paragraph_index(doc, heading)
    end = len(paragraphs)
    for i in range(start + 1, len(paragraphs)):
        if paragraphs[i].text.strip() == next_heading_text:
            end = i
            break
    for paragraph in list(paragraphs[start + 1:end]):
        if not paragraph.text.strip():
            delete_paragraph(paragraph)
    anchor = heading
    for item in items:
        anchor = insert_paragraph_after(anchor, item, "Normal")


def main():
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(str(OUTPUT))

    # Replace obvious placeholders with content consistent with the four models.
    try:
        replace_paragraph(find_paragraph(doc, "基于问题研究", exact=True), "药材烘干过程的热湿耦合与移动边界研究")
    except ValueError:
        pass

    # The source uses the heading text “摘 要” with an internal space.
    abstract_heading = find_paragraph(doc, "摘", exact=False)
    abstract = (
        "针对中药材热风烘干过程中温度场与水分场时空分布难以直接获得、尺寸变化又会影响传递过程的问题，"
        "本文以近似圆柱形药材为研究对象，建立固定边界与移动边界条件下的径向传热传质模型。首先，采用初值固定的 "
        "stretched-exponential（拉伸指数）函数对烘房环境数据进行连续化拟合；随后基于守恒定律建立温度和水分浓度控制方程，"
        "并采用有限体积法、Kirchhoff 界面平均和带稀疏解析 Jacobian 的隐式 BDF 方法求解。针对全过程烘干，进一步引入随温度和含水率变化的物性参数；"
        "针对尺寸变化，利用 PCHIP 拟合半径并通过材料坐标将移动区域映射到固定区域。结果表明，预热平衡阶段药材表面升温和失水均快于中心，"
        "1800 s 时中心与表面温度分别为 33.5689 ℃ 和 36.7806 ℃，表面含水率降至 1.5102 kg/kg；固定半径模型下药材达到含水率要求的时间为 57.6785 h，"
        "考虑收缩后为 51.2682 h，最终半径约为 1.2000 cm。网格加密、时间精度和守恒检验表明，数值结果具有良好的稳定性。"
        "本文模型为中药材干燥过程的质量判定和工艺优化提供了定量依据。"
    )
    abstract_paragraph = None
    for paragraph in doc.paragraphs:
        if paragraph._p is abstract_heading._p:
            continue
        if paragraph.text.strip() == "随":
            abstract_paragraph = paragraph
            break
    if abstract_paragraph is None:
        abstract_paragraph = insert_paragraph_after(abstract_heading)
    replace_paragraph(abstract_paragraph, abstract)
    abstract_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    keyword_text = "关键词：药材烘干；热湿耦合；有限体积法；隐式BDF；移动边界"
    if not any(p.text.strip() == keyword_text for p in doc.paragraphs):
        keyword_paragraph = insert_paragraph_after(
            abstract_paragraph,
            keyword_text,
            "Normal",
        )
        keyword_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    assumptions = [
        "将药材视为均匀、各向同性的近似圆柱体，长度方向和周向上的温度、水分浓度及物性差异忽略，仅考虑中心对称条件下的径向传递。",
        "问题一至问题三中药材半径保持不变；问题四中假设长度基本不变、药材沿径向均匀收缩，材料点的相对位置保持不变，且不考虑开裂、翘曲等局部变形。",
        "药材内部传热遵循傅里叶定律，水分迁移采用等效扩散模型；忽略蒸发潜热、湿分迁移引起的显热交换、内部热源及形变功等二阶影响。",
        "药材与烘房之间通过有限速率的对流换热和对流传质交换能量与水分，表面满足 Robin 边界条件；附件环境数据用连续函数拟合，分界与外推过程保持连续。",
        "模型参数及经验物性关系在题设温度、含水率和时间范围内有效，测量与拟合误差通过网格收敛、时间精度、守恒性和物理约束检验进行控制。",
    ]
    assumption_heading = find_paragraph(doc, "模型假设", exact=True)
    symbol_heading = find_paragraph(doc, "符号说明", exact=True)
    paragraphs = doc.paragraphs
    start = paragraph_index(doc, assumption_heading)
    end = paragraph_index(doc, symbol_heading)
    body = [p for p in paragraphs[start + 1:end] if p.text.strip()]
    for paragraph, text in zip(body, assumptions):
        replace_paragraph(paragraph, text)
    if len(body) < len(assumptions):
        anchor = body[-1] if body else assumption_heading
        for text in assumptions[len(body):]:
            anchor = insert_paragraph_after(anchor, text, "Normal")

    symbol_rows = [
        ("t", "时间", "s"),
        ("r", "距药材中心的径向坐标", "m"),
        ("R", "药材初始半径", "m"),
        ("R(t)", "t 时刻药材半径", "m"),
        ("L", "药材长度", "m"),
        ("ξ", "材料坐标，ξ=r/R(t)", "—"),
        ("T(r,t)", "t 时刻位置 r 处的温度", "℃"),
        ("C(r,t)", "t 时刻位置 r 处的干基水分浓度", "kg/kg"),
        ("Θ(ξ,t)", "材料坐标下的温度场", "℃"),
        ("U(ξ,t)", "材料坐标下的水分浓度场", "kg/kg"),
        ("T_a(t)", "烘房空气温度", "℃"),
        ("C_a(t)", "烘房空气水分浓度", "kg/kg"),
        ("ρ(C)", "与水分浓度相关的密度", "kg/m³"),
        ("c_p(C)", "比热容", "J/(kg·K)"),
        ("k(C)", "热传导系数", "W/(m·K)"),
        ("h", "对流换热系数", "W/(m²·K)"),
        ("h_m", "对流传质系数", "m/s"),
        ("D(C), D(C,T)", "有效水分扩散系数", "m²/s"),
        ("α", "热扩散率，α=k/(ρc_p)", "m²/s"),
        ("v", "边界收缩速度，v=dR/dt", "m/s"),
        ("M(t)", "全域最大水分浓度", "kg/kg"),
        ("C_cr", "干燥达标临界水分浓度", "kg/kg"),
        ("t_*", "首次满足 M(t)<C_cr 的时刻", "s 或 h"),
    ]
    table = doc.tables[0]
    while len(table.rows) < len(symbol_rows) + 1:
        table.add_row()
    set_cell(table.rows[0].cells[0], "符号", 9)
    set_cell(table.rows[0].cells[1], "含义", 9)
    set_cell(table.rows[0].cells[2], "单位", 9)
    for row, values in zip(table.rows[1:], symbol_rows):
        for cell, value in zip(row.cells, values):
            set_cell(cell, value, 8.5)

    fill_section(
        doc,
        "模型的优点",
        "模型的不足",
        [
            "（1）模型具有明确的物理含义：以圆柱坐标下的热传导和水分扩散方程为核心，并通过对流边界刻画药材与烘房环境的交换过程，能够给出内部任意时空位置的温度与水分浓度。",
            "（2）环境边界和半径数据均经过连续化与物理约束处理，拉伸指数拟合、分段平滑过渡及 PCHIP 插值避免了数据突变、单调性破坏和半径过拟合。",
            "（3）有限体积法保持守恒，Kirchhoff 界面平均适用于非线性扩散系数；隐式 BDF 与稀疏解析 Jacobian 能够稳定处理长时间、多尺度和强非线性的计算。",
            "（4）模型同时覆盖固定边界和收缩边界，并通过达标事件检测直接给出干燥时间，具有较好的结果可解释性和工程使用价值。",
        ],
    )
    fill_section(
        doc,
        "模型的不足",
        "模型的推广",
        [
            "（1）当前模型主要考虑径向传递，忽略了端部效应、轴向不均匀性以及药材内部孔隙结构的空间差异；对于细长比不足或形状不规则的药材，近似圆柱假设会带来误差。",
            "（2）模型将传热与传质的耦合影响进行了适度简化，未显式考虑蒸发潜热、湿分迁移显热和局部相变等过程；物性关系和对流系数也依赖附件数据及经验参数。",
            "（3）问题四采用长度不变、径向均匀收缩假设，尚未描述开裂、翘曲、分层等复杂形变；此外，模型尚未进一步优化能耗、空气流量等实际工艺控制变量。",
        ],
    )
    fill_section(
        doc,
        "模型的推广",
        "AI工具使用声明",
        [
            "（1）在能量方程中加入蒸发潜热、湿分迁移显热和局部相变项，并令扩散系数、换热系数及传质系数同时依赖温度、含水率和结构状态，以建立更完整的热湿耦合模型。",
            "（2）将一维径向模型推广至二维或三维，并结合动网格、有限元或 ALE 方法描述长度变化、各向异性收缩、裂纹和翘曲等复杂形变。",
            "（3）利用更多批次实验数据进行反问题辨识和贝叶斯校准，量化参数、边界测量及拟合误差对干燥时间和质量指标的影响。",
            "（4）在数值模型基础上引入最优控制或模型预测控制，以干燥时间、能耗和成品均匀性为多目标，优化温度、湿度和风速的动态调节策略。",
        ],
    )

    flowcharts = {
        "问题一模型的建立与求解": (
            "问题一总体思路",
            [
                "题目数据与初始条件",
                "边界数据连续化\n拉伸指数模型\nstretched-\nexponential",
                "固定半径一维径向模型\n热传导方程 + 水分扩散方程",
                "有限体积法离散\nKirchhoff 界面平均",
                "隐式 BDF 求解\n稀疏解析 Jacobian",
                "网格与时间\n收敛性检验",
                "输出温度/水分\n时空分布与结果表",
            ],
        ),
        "问题二模型的建立与求解": (
            "问题二总体思路",
            [
                "附件全时段\n环境数据",
                "SG 平滑与稳定分界\n温度/水分分别识别",
                "分段拟合与平滑过渡\n构造连续边界",
                "变物性热湿耦合\n固定半径 PDE 模型",
                "有限体积法 + 隐式 BDF\n四分块稀疏\nJacobian",
                "前 3 h 数值计算\n误差与守恒检验",
                "温度与含水率\n场分布结果",
            ],
        ),
        "问题三模型的建立与求解": (
            "问题三总体思路",
            [
                "沿用问题二\n热湿耦合模型",
                "尾段边界外推\n延长计算时间",
                "表面加密径向网格\n提高阈值捕捉精度",
                "计算全域最大值\nM(t)=max C(r,t)",
                "事件检测\nM(t)<0.15",
                "网格、时间与\n质量守恒校验",
                "确定达标时间\n输出各时刻结果",
            ],
        ),
        "问题四模型的建立与求解": (
            "问题四总体思路",
            [
                "附件半径离散数据",
                "物理约束筛选\nPCHIP 拟合 R(t)",
                "建立移动边界\n0≤r≤R(t)",
                "材料坐标变换\nξ=r/R(t) → [0,1]",
                "变物性移动边界\n热湿耦合方程",
                "有限体积法 + BDF\n固定区域求解",
                "映射回实际距离\n全域达标判定",
            ],
        ),
    }
    for heading_text, (title, nodes) in flowcharts.items():
        image_path = FLOW_DIR / ("flow_" + heading_text[:3] + ".png")
        draw_flowchart(image_path, title, nodes)
        anchor = find_paragraph(doc, heading_text, exact=True)
        number = {"问题一": "11", "问题二": "12", "问题三": "13", "问题四": "14"}[heading_text[:3]]
        add_flowchart(doc, anchor, image_path, f"图 {number} {title}流程图")

    doc.save(str(OUTPUT))
    print(OUTPUT)
    for path in sorted(FLOW_DIR.glob("flow_*.png")):
        print(path)


if __name__ == "__main__":
    main()
