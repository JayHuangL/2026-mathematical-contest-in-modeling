from pathlib import Path
import shutil

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\Code\2026-mathematical-contest-in-modeling")
SOURCE = ROOT / "paper" / "2026C论文.docx"
OUTPUT = ROOT / "paper" / "2026C论文_参照优秀论文修订版.docx"
FLOW_DIR = ROOT / "paper" / "tmp" / "20260913_honors_flowcharts"
FLOW_DIR.mkdir(parents=True, exist_ok=True)


def normalize(text):
    return "".join(text.split())


def find_paragraph(doc, text, exact=False):
    target = normalize(text)
    for paragraph in doc.paragraphs:
        value = normalize(paragraph.text)
        if value == target if exact else target in value:
            return paragraph
    raise ValueError(f"找不到段落：{text}")


def paragraph_index(doc, target):
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph._p is target._p:
            return i
    raise ValueError("目标段落不在文档中")


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def remove_num_pr(paragraph):
    ppr = paragraph._p.pPr
    if ppr is not None:
        num_pr = ppr.find(qn("w:numPr"))
        if num_pr is not None:
            ppr.remove(num_pr)


def replace_paragraph(paragraph, text, style=None):
    clear_paragraph(paragraph)
    remove_num_pr(paragraph)
    if style:
        paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def insert_paragraph_after(paragraph, text=None, style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_paragraph.style = style
        except KeyError:
            pass
    if text is not None:
        new_paragraph.add_run(text)
    return new_paragraph


def remove_paragraph(paragraph):
    paragraph._element.getparent().remove(paragraph._element)


def add_rich_text(paragraph, label, body, bold_phrases=(), justify=True):
    clear_paragraph(paragraph)
    remove_num_pr(paragraph)
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    label_run = paragraph.add_run(label)
    label_run.bold = True
    remaining = body
    while remaining:
        match = None
        for phrase in bold_phrases:
            pos = remaining.find(phrase)
            if pos >= 0 and (match is None or pos < match[0]):
                match = (pos, phrase)
        if match is None:
            paragraph.add_run(remaining)
            break
        pos, phrase = match
        if pos:
            paragraph.add_run(remaining[:pos])
        run = paragraph.add_run(phrase)
        run.bold = True
        remaining = remaining[pos + len(phrase):]
    return paragraph


def set_cell(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.bold = bold


def fill_body_section(doc, heading_text, next_heading_text, paragraphs_text):
    heading = find_paragraph(doc, heading_text, exact=True)
    start = paragraph_index(doc, heading)
    doc_paragraphs = doc.paragraphs
    end = len(doc_paragraphs)
    target_next = normalize(next_heading_text)
    for i in range(start + 1, len(doc_paragraphs)):
        if normalize(doc_paragraphs[i].text) == target_next:
            end = i
            break
    for paragraph in list(doc_paragraphs[start + 1:end]):
        if not paragraph.text.strip():
            remove_paragraph(paragraph)
    anchor = heading
    for text in paragraphs_text:
        anchor = insert_paragraph_after(anchor, text, "Normal")
        anchor.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def choose_font(size):
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
        Path(r"C:\Windows\Fonts\Deng.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                return ImageFont.truetype(str(candidate), size, index=0)
            except Exception:
                try:
                    return ImageFont.truetype(str(candidate), size)
                except Exception:
                    pass
    return ImageFont.load_default()


def wrap_text(draw, text, font, max_width):
    lines = []
    for part in text.split("\n"):
        current = ""
        for char in part:
            candidate = current + char
            bbox = draw.textbbox((0, 0), candidate, font=font)
            if current and bbox[2] - bbox[0] > max_width:
                lines.append(current)
                current = char
            else:
                current = candidate
        if current:
            lines.append(current)
    return lines


def draw_flowchart(path, nodes):
    width, height = 2200, 480
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    node_font = choose_font(29)
    node_width = 300
    node_height = 175
    margin = 45
    gap = (width - 2 * margin - len(nodes) * node_width) / (len(nodes) - 1)
    top = 150
    outline = (42, 56, 76)
    fill = (247, 249, 252)
    positions = []

    for i, node in enumerate(nodes):
        x = int(margin + i * (node_width + gap))
        y = top
        positions.append((x, y))
        draw.rectangle((x, y, x + node_width, y + node_height), fill=fill, outline=outline, width=3)
        lines = wrap_text(draw, node, node_font, node_width - 32)
        heights = []
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=node_font)
            heights.append(bbox[3] - bbox[1])
        total = sum(heights) + 8 * max(0, len(lines) - 1)
        current_y = y + (node_height - total) / 2
        for line, line_height in zip(lines, heights):
            bbox = draw.textbbox((0, 0), line, font=node_font)
            text_width = bbox[2] - bbox[0]
            draw.text((x + (node_width - text_width) / 2, current_y), line, fill=outline, font=node_font)
            current_y += line_height + 8

    arrow_y = top + node_height / 2
    for i in range(len(positions) - 1):
        x1 = positions[i][0] + node_width + 7
        x2 = positions[i + 1][0] - 9
        draw.line((x1, arrow_y, x2, arrow_y), fill=outline, width=4)
        draw.polygon([(x2, arrow_y), (x2 - 15, arrow_y - 10), (x2 - 15, arrow_y + 10)], fill=outline)
    image.save(path, dpi=(200, 200))


def insert_flowchart(doc, heading_text, image_path, caption):
    anchor = find_paragraph(doc, heading_text, exact=True)
    if heading_text.startswith("问题四"):
        anchor.paragraph_format.page_break_before = True
    image_paragraph = insert_paragraph_after(anchor)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.15))
    caption_paragraph = insert_paragraph_after(image_paragraph)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        caption_paragraph.style = "图注"
    except KeyError:
        caption_paragraph.style = "Normal"
    run = caption_paragraph.add_run(caption)
    run.font.size = Pt(9)


def main():
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(str(OUTPUT))

    # The original title is a placeholder; use the subject of the submitted paper.
    try:
        replace_paragraph(
            find_paragraph(doc, "基于问题研究", exact=True),
            "药材烘干过程的热湿耦合与移动边界研究",
        )
    except ValueError:
        pass

    abstract_heading = find_paragraph(doc, "摘要", exact=False)
    heading_idx = paragraph_index(doc, abstract_heading)
    abstract_body = doc.paragraphs[heading_idx + 1]
    intro = (
        "热风烘干是中药材加工过程中决定成品质量的重要环节，温度与水分浓度在药材内部的时空分布以及干燥过程中的尺寸收缩，"
        "直接影响干燥效率和成品品质。本文针对药材热风烘干过程中的热湿传递问题，建立数学模型，分析药材内部温度场与水分场的演化规律，"
        "并确定达到干燥要求所需的时间。"
    )
    replace_paragraph(abstract_body, intro, "Normal")
    abstract_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    abstract_items = [
        (
            "针对问题一，",
            "将药材简化为固定半径的轴对称圆柱体，根据能量守恒和水分质量守恒分别建立温度传导模型和水分扩散模型；"
            "对附件1环境数据采用拉伸指数模型（stretched-exponential）进行连续拟合，并采用有限体积法和隐式BDF方法求解，"
            "得到预热平衡阶段各时刻、各径向位置的温度和水分浓度。结果表明，1800 s时药材中心与表面温度分别为33.5689 ℃和36.7806 ℃，"
            "表面水分浓度降至1.5102 kg/kg。",
            ("有限体积法和隐式BDF方法", "33.5689 ℃和36.7806 ℃", "1.5102 kg/kg"),
        ),
        (
            "针对问题二，",
            "考虑预热平衡与恒温干燥两个阶段，先利用Savitzky-Golay平滑和连续窗口判别温度、水分浓度的稳定分界点，"
            "再用分段拉伸指数函数平滑连接两阶段边界，并将物性参数表示为水分浓度或温度、水分浓度的函数，建立变物性热湿耦合模型。"
            "通过有限体积法与隐式BDF联立求解，得到0～3 h内的场分布；3 h时中心与表面温度差约为0.0421 ℃，水分浓度分别为1.7658 kg/kg和1.0075 kg/kg。",
            ("变物性热湿耦合模型", "0.0421 ℃", "1.7658 kg/kg和1.0075 kg/kg"),
        ),
        (
            "针对问题三，",
            "沿用问题二的热湿耦合模型，以药材全域最大水分浓度首次低于0.15 kg/kg作为干燥完成判据，"
            "采用表面加密网格和事件检测确定临界时刻。计算得到固定半径条件下的干燥时间为57.6785 h，此时中心水分浓度约为0.1500 kg/kg。",
            ("57.6785 h", "0.1500 kg/kg"),
        ),
        (
            "针对问题四，",
            "先对附件2的半径时序数据进行物理约束筛选和PCHIP插值，得到连续半径函数；再引入材料坐标，将随时间变化的物理区域映射到固定区间，"
            "建立考虑收缩的移动边界热湿耦合模型。结果表明，考虑药材收缩后达到同一干燥要求的时间缩短为51.2682 h，末期半径约为1.2000 cm。"
            "模型通过网格收敛、时间精度和守恒性检验，具有较好的稳定性，可为药材干燥工艺设计提供参考。",
            ("PCHIP插值", "51.2682 h", "1.2000 cm"),
        ),
    ]
    anchor = abstract_body
    for label, body_text, bold_phrases in abstract_items:
        paragraph = insert_paragraph_after(anchor, style="Normal")
        add_rich_text(paragraph, label, body_text, bold_phrases)
        anchor = paragraph

    keyword_text = "关键词：药材烘干；热湿耦合；有限体积法；隐式BDF；移动边界"
    keyword_paragraphs = [p for p in doc.paragraphs if normalize(p.text).startswith(normalize("关键词"))]
    if keyword_paragraphs:
        replace_paragraph(keyword_paragraphs[0], keyword_text, "Normal")
        keyword_paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        keyword = insert_paragraph_after(anchor, keyword_text, "Normal")
        keyword.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Keep the abstract and keywords on the cover page, as in the reference
    # award-winning papers; start the main text on the following page.
    find_paragraph(doc, "问题重述", exact=True).paragraph_format.page_break_before = True

    assumptions = [
        "1、假设药材为均匀、各向同性的近似圆柱体，忽略轴向和周向差异，仅考虑径向传热传质。",
        "2、问题一至问题三中药材半径保持不变；问题四中假设长度不变且沿径向均匀收缩，忽略开裂、翘曲等局部变形。",
        "3、传热遵循傅里叶定律，水分迁移采用等效扩散定律；忽略蒸发潜热、湿分迁移显热及内部热源等影响。",
        "4、药材与烘房通过对流换热和对流传质交换热量与水分，表面满足Robin边界条件，环境数据用连续函数拟合。",
        "5、题设给出的物性关系、换热系数和传质系数在研究时间及含水率范围内有效。",
    ]
    assumption_heading = find_paragraph(doc, "模型假设", exact=True)
    symbol_heading = find_paragraph(doc, "符号说明", exact=True)
    start = paragraph_index(doc, assumption_heading)
    end = paragraph_index(doc, symbol_heading)
    body_paragraphs = [p for p in doc.paragraphs[start + 1:end] if p.text.strip()]
    for paragraph, text in zip(body_paragraphs, assumptions):
        replace_paragraph(paragraph, text, "Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if len(body_paragraphs) < len(assumptions):
        insert_anchor = body_paragraphs[-1] if body_paragraphs else assumption_heading
        for text in assumptions[len(body_paragraphs):]:
            insert_anchor = insert_paragraph_after(insert_anchor, text, "Normal")

    table = doc.tables[0]
    if len(table.columns) == 3:
        table.add_column(Inches(2.0))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    while len(table.rows) < 13:
        table.add_row()
    while len(table.rows) > 13:
        table._tbl.remove(table.rows[-1]._tr)
    widths = [Inches(0.75), Inches(2.55), Inches(0.95), Inches(2.25)]
    for column, width in zip(table.columns, widths):
        column.width = width
    headers = ["符号", "含义", "符号", "含义"]
    for j, header in enumerate(headers):
        set_cell(table.rows[0].cells[j], header, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    symbol_pairs = [
        ("t", "时间（s）"),
        ("r", "径向坐标（m）"),
        ("R", "初始半径（m）"),
        ("R(t)", "t时刻半径（m）"),
        ("L", "药材长度（m）"),
        ("ξ", "材料坐标，ξ=r/R(t)"),
        ("T(r,t)", "位置r处温度（℃）"),
        ("C(r,t)", "位置r处水分浓度（kg/kg）"),
        ("T_a(t)", "烘房空气温度（℃）"),
        ("C_a(t)", "烘房空气水分浓度（kg/kg）"),
        ("ρ(C)", "密度（kg/m³）"),
        ("c_p(C)", "比热容（J/(kg·K)）"),
        ("k(C)", "热传导系数（W/(m·K)）"),
        ("h", "对流换热系数（W/(m²·K)）"),
        ("h_m", "对流传质系数（m/s）"),
        ("D(C)", "问题一有效扩散系数（m²/s）"),
        ("D(C,T)", "问题二至四有效扩散系数（m²/s）"),
        ("α", "热扩散率，α=k/(ρc_p)"),
        ("v", "边界收缩速度，v=dR/dt（m/s）"),
        ("M(t)", "全域最大水分浓度（kg/kg）"),
        ("C_cr", "干燥达标临界水分浓度（kg/kg）"),
        ("t_*", "首次满足M(t)<C_cr的时刻（s或h）"),
        ("Θ(ξ,t)", "材料坐标下温度场（℃）"),
        ("U(ξ,t)", "材料坐标下水分浓度场（kg/kg）"),
    ]
    for i in range(12):
        left = symbol_pairs[i]
        right = symbol_pairs[i + 12]
        set_cell(table.rows[i + 1].cells[0], left[0], size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(table.rows[i + 1].cells[1], left[1], size=8.5)
        set_cell(table.rows[i + 1].cells[2], right[0], size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(table.rows[i + 1].cells[3], right[1], size=8.5)

    fill_body_section(
        doc,
        "模型的优点",
        "模型的不足",
        [
            "本文以能量守恒和水分质量守恒为基础，建立了与药材圆柱几何和烘房边界相对应的热湿传递模型，能够给出药材内部不同位置、不同时间的温度和水分浓度。",
            "针对附件数据的离散性，分别采用拉伸指数函数和PCHIP方法构造连续边界与半径函数，并通过单调性、越界性和拟合误差进行筛选，使输入条件具有较好的物理合理性。",
            "有限体积法保证离散过程的守恒性，Kirchhoff界面平均适用于非线性扩散系数，隐式BDF方法能够稳定求解长时间、多尺度的热湿耦合问题；同时，网格、时间步长及质量守恒检验为计算结果提供了可靠性依据。",
        ],
    )
    fill_body_section(
        doc,
        "模型的不足",
        "模型的推广",
        [
            "模型将药材视为轴对称圆柱体，忽略端部效应、轴向不均匀性和孔隙结构差异，对于形状不规则或内部组织差异较大的药材，计算结果可能存在偏差。",
            "模型对蒸发潜热、湿分迁移显热、局部相变及复杂形变进行了简化处理，且物性关系和边界传递系数依赖题设经验公式，长时间外推时仍存在参数不确定性。",
            "本文主要研究温度、水分浓度和达标时间，尚未把能耗、风速和成品均匀性等工艺指标纳入统一优化框架。",
        ],
    )
    fill_body_section(
        doc,
        "模型的推广",
        "AI工具使用声明",
        [
            "可在能量方程中加入蒸发潜热、湿分迁移显热和局部相变项，并令扩散系数、换热系数和传质系数同时依赖温度、含水率及结构状态，建立更加完整的热湿耦合模型。",
            "可将一维径向模型推广至二维或三维，并结合动网格、有限元或ALE方法描述长度变化、各向异性收缩、裂纹和翘曲等复杂形变。",
            "可利用多批次实验数据进行反问题辨识和贝叶斯校准，量化物性参数、环境边界及测量误差对干燥时间和成品均匀性的影响。",
            "可在数值模型基础上引入最优控制或模型预测控制，以干燥时间、能耗和成品均匀性为目标，优化温度、湿度和风速的动态调节策略。",
        ],
    )

    # Remove source-draft editorial marks from the submission-style copy.
    try:
        replace_paragraph(
            find_paragraph(doc, "建立材料坐标下的水分守恒关系", exact=False),
            "建立材料坐标下的水分守恒关系",
        )
    except ValueError:
        pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = None

    charts = {
        "问题一模型的建立与求解": [
            "附件1、2数据",
            "连续拟合边界",
            "固定半径径向模型",
            "有限体积法离散",
            "隐式BDF求解",
            "输出温度与水分场",
        ],
        "问题二模型的建立与求解": [
            "附件1全时段数据",
            "SG平滑识别分界",
            "分段函数平滑连接",
            "变物性热湿耦合模型",
            "有限体积法+BDF",
            "输出0～3 h结果",
        ],
        "问题三模型的建立与求解": [
            "问题二模型",
            "延长边界与时间",
            "表面加密网格",
            "计算全域最大值M(t)",
            "事件检测M(t)<0.15",
            "确定干燥时间",
        ],
        "问题四模型的建立与求解": [
            "附件2半径数据",
            "物理约束筛选",
            "PCHIP拟合R(t)",
            "材料坐标变换",
            "移动边界热湿模型",
            "输出实际位置结果",
        ],
    }
    for heading_text, nodes in charts.items():
        number = {"问题一": "11", "问题二": "12", "问题三": "13", "问题四": "14"}[heading_text[:3]]
        image_path = FLOW_DIR / ("flow_" + heading_text[:3] + ".png")
        draw_flowchart(image_path, nodes)
        insert_flowchart(doc, heading_text, image_path, f"图 {number} {heading_text[:3]}流程图")

    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
