from pathlib import Path
import argparse,json,hashlib
import sys
import numpy as np
import openpyxl
HERE=Path(__file__).resolve().parent
PACKAGE=HERE.parent.parent
if str(HERE.parent) not in sys.path:
    sys.path.insert(0,str(HERE.parent))
from artifact_names import artifact_name, workbook_path
p=argparse.ArgumentParser()
p.add_argument('--root',type=Path,default=PACKAGE/'data')
p.add_argument('--book',type=Path,default=workbook_path(PACKAGE/'results','q3'))
a=p.parse_args()
book=a.book
s=json.loads((PACKAGE/'results'/'q3'/artifact_name('q3','validation')).read_text(encoding='utf-8'))
payload=json.loads((PACKAGE/'results'/'q3'/artifact_name('q3','result_data')).read_text(encoding='utf-8'))
assert hashlib.sha256((PACKAGE/'data'/'附件1.xlsx').read_bytes()).hexdigest()==s['input_sha256']
assert hashlib.sha256((PACKAGE/'code'/'q2'/'solve_q2.py').read_bytes()).hexdigest()==s['q2_code_sha256']
with np.load(PACKAGE/'results'/'q3'/artifact_name('q3','full_precision')) as full:
    expected=np.round(full['C'][1:],4)
    assert np.array_equal(expected,payload['C'])
    assert np.array_equal(np.round(full['t'][1:],4),payload['t_s'])
    assert np.all(np.diff(full['t'][:-1])==60)
    assert full['C'][-1].max()<.15 and full['C'][-2].max()>.15
    assert full['t'][-2]<s['event_s']<full['t'][-1]
    for t,c in zip(s['table_t_s'],s['table_C']):
        i=int(np.argmin(abs(full['t']-t)))
        assert np.array_equal(full['C'][i,::5],c)
w=openpyxl.load_workbook(book,data_only=True,read_only=True)
assert w.sheetnames==['Sheet1']
rows=iter(w.active.iter_rows())
header=next(rows)
assert [c.value for c in header[1:]]==payload['r_cm']
count=0
for i,row in enumerate(rows):
    assert row[0].value==payload['t_s'][i]
    assert np.array_equal([c.value for c in row[1:]],expected[i])
    assert all(c.number_format=='0.0000' for c in row[1:])
    count+=1
assert count==s['output_time_rows']
w.close()
result={'time_rows':count,'numeric_cells':count*21,'all_values_match':True,'strict_terminal_check':True,'input_and_q2_unchanged':True}
(PACKAGE/'results'/'q3'/artifact_name('q3','validation_export')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
