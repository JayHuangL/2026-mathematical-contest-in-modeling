"""根据 submit/results 中的 JSON 结果载荷构建四个结果工作簿。"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import openpyxl

from artifact_names import artifact_name, workbook_path


HERE = Path(__file__).resolve().parent
PACKAGE = HERE.parent
DATA = PACKAGE / "data"
RESULTS = PACKAGE / "results"


def _clear_rows(sheet) -> None:
    if sheet.max_row and sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)


def _write_two_sheet(question: str, template_name: str, keys: tuple[str, str]) -> None:
    folder = RESULTS / question
    payload = json.loads((folder / artifact_name(question, "result_data")).read_text(encoding="utf-8"))
    book = openpyxl.load_workbook(DATA / template_name)
    if book.sheetnames[:2] != ["温度", "水分浓度"]:
        raise ValueError(f"Unexpected template sheets for {question}: {book.sheetnames}")
    for sheet_name, key in zip(("温度", "水分浓度"), keys):
        sheet = book[sheet_name]
        sheet.cell(1, 1, "时间\\到药材中心的距离")
        for col, distance in enumerate(payload["r_cm"], 2):
            sheet.cell(1, col, distance).number_format = "0.0"
        for row, (time_s, values) in enumerate(zip(payload["t_s"], payload[key]), 2):
            sheet.cell(row, 1, time_s).number_format = "0"
            for col, value in enumerate(values, 2):
                sheet.cell(row, col, value).number_format = "0.0000"
        sheet.freeze_panes = "B2"
    book.save(workbook_path(RESULTS, question))
    book.close()


def _write_q3() -> None:
    folder = RESULTS / "q3"
    payload = json.loads((folder / artifact_name("q3", "result_data")).read_text(encoding="utf-8"))
    output = workbook_path(RESULTS, "q3")
    shutil.copy2(DATA / "result3_template.xlsx", output)
    book = openpyxl.load_workbook(output)
    sheet = book["Sheet1"]
    _clear_rows(sheet)
    headers = ["时间\\到药材中心的距离", *payload["r_cm"]]
    for col, value in enumerate(headers, 1):
        sheet.cell(1, col, value)
    for row, (time_s, values) in enumerate(zip(payload["t_s"], payload["C"]), 2):
        sheet.cell(row, 1, time_s).number_format = "0.0000" if row == len(payload["t_s"]) + 1 else "0"
        for col, value in enumerate(values, 2):
            sheet.cell(row, col, value).number_format = "0.0000"
    sheet.freeze_panes = "B2"
    book.save(output)
    book.close()


def _write_q4() -> None:
    folder = RESULTS / "q4"
    payload = json.loads((folder / artifact_name("q4", "result_data")).read_text(encoding="utf-8"))
    output = workbook_path(RESULTS, "q4")
    shutil.copy2(DATA / "result4_template.xlsx", output)
    book = openpyxl.load_workbook(output)
    sheet = book["Sheet1"]
    _clear_rows(sheet)
    headers = ["时间\\到药材中心的距离", *payload["fixed_r_cm"], "药材表面"]
    for col, value in enumerate(headers, 1):
        sheet.cell(1, col, value)
    for row, (time_s, values) in enumerate(zip(payload["t_s"], payload["C"]), 2):
        sheet.cell(row, 1, time_s).number_format = "0.0000" if row == len(payload["t_s"]) + 1 else "0"
        for col, value in enumerate(values, 2):
            sheet.cell(row, col, value).number_format = "0.0000"
    sheet.freeze_panes = "B2"
    book.save(output)
    book.close()


def main() -> None:
    _write_two_sheet("q1", "result1_template.xlsx", ("T", "C"))
    _write_two_sheet("q2", "result2_template.xlsx", ("T", "C"))
    _write_q3()
    _write_q4()
    print("rebuilt result1.xlsx, result2.xlsx, result3.xlsx and result4.xlsx")


if __name__ == "__main__":
    main()
