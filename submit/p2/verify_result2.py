"""Read-only verification of exported workbook against computed full-precision data."""
from pathlib import Path
import hashlib
import json
import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
summary = json.loads((HERE/'validation.json').read_text(encoding='utf-8'))
assert hashlib.sha256((HERE/'data'/'附件1.xlsx').read_bytes()).hexdigest() == summary['input_sha256']
data = json.loads((HERE/'result2_data.json').read_text(encoding='utf-8'))
book = openpyxl.load_workbook(HERE/'result2.xlsx', read_only=True, data_only=True)
assert book.sheetnames == ['温度', '水分浓度']
checks = []
with np.load(HERE/'result2_full_precision.npz') as full:
    for name, key in [('温度', 'T'), ('水分浓度', 'C')]:
        expected = np.round(full[key][1:], 4)
        assert np.array_equal(expected, np.array(data[key]))
        rows = iter(book[name].iter_rows())
        header = next(rows)
        assert len(header) == 22
        assert [cell.value for cell in header[1:]] == data['r_cm']
        count = 0
        for i, row in enumerate(rows):
            assert row[0].value == i+1 == data['t_s'][i]
            assert len(row) == 22
            assert np.array_equal(np.array([cell.value for cell in row[1:]]), expected[i])
            assert all(cell.number_format == '0.0000' for cell in row[1:])
            count += 1
        assert count == 10800
        checks.append({'sheet': name, 'time_rows': count, 'distance_columns': 21,
                       'numeric_values_verified': count*21})
book.close()
report = (HERE/'p2.md').read_text(encoding='utf-8')
tables = (HERE/'结果表.md').read_text(encoding='utf-8')
assert '{{' not in report
assert all(line in report for line in tables.splitlines() if line.startswith('|'))
for key in ['T', 'C']:
    expected_table = np.round(np.array(summary[f'table_{key}']), 4)
    with np.load(HERE/'result2_full_precision.npz') as full:
        assert np.array_equal(expected_table, np.round(full[key][1800::1800, ::5], 4))
result = {'checks': checks, 'all_453600_values_match': True,
          'markdown_tables_match': True, 'input_unchanged': True}
(HERE/'validation_export.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False))
