"""Lightweight independent checks for the integrated submission package."""
from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

import numpy as np
import openpyxl


HERE = Path(__file__).resolve().parent
PACKAGE = HERE.parent
DATA = PACKAGE / "data"
RESULTS = PACKAGE / "results"


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _csv_records(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def _payload_check(question: str, expected_rows: int, keys: tuple[str, ...]) -> dict:
    folder = RESULTS / question
    payload = json.loads((folder / f"result{question[1:]}_data.json").read_text(encoding="utf-8"))
    times = payload["t_s"]
    if len(times) != expected_rows:
        raise AssertionError(f"{question}: wrong output row count")
    for key in keys:
        values = payload[key]
        if len(values) != expected_rows or any(len(row) != 21 for row in values):
            raise AssertionError(f"{question}: wrong {key} shape")
    return {"rows": len(times), "radial_columns": 21, "payload": True}


def _workbook_shape(path: Path, sheets: list[str], rows: int, columns: int) -> dict:
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if book.sheetnames[: len(sheets)] != sheets:
        raise AssertionError(f"{path}: unexpected sheets {book.sheetnames}")
    counts = []
    for sheet in book.worksheets[: len(sheets)]:
        count = sum(1 for _ in sheet.iter_rows(min_row=2, max_col=columns))
        counts.append(count)
    book.close()
    if counts != [rows] * len(sheets):
        raise AssertionError(f"{path}: unexpected rows {counts}")
    return {"sheets": sheets, "rows": counts}


def _assert_cell(actual, expected, label: str) -> None:
    if expected is None:
        if actual is not None:
            raise AssertionError(f"{label}: expected blank cell")
    elif isinstance(expected, str):
        if actual != expected:
            raise AssertionError(f"{label}: workbook header differs from payload")
    elif actual is None or not np.isclose(float(actual), float(expected), rtol=0.0, atol=5e-12):
        raise AssertionError(f"{label}: workbook value differs from payload")


def _workbook_matches_payload(question: str, sheet_keys: list[tuple[str, str]], headers: list[object]) -> dict:
    folder = RESULTS / question
    payload = json.loads((folder / f"result{question[1:]}_data.json").read_text(encoding="utf-8"))
    book = openpyxl.load_workbook(folder / f"result{question[1:]}.xlsx", data_only=True, read_only=True)
    if book.sheetnames[:len(sheet_keys)] != [name for name, _ in sheet_keys]:
        raise AssertionError(f"{question}: unexpected workbook sheets")
    times = payload["t_s"]
    for sheet_name, key in sheet_keys:
        rows = iter(book[sheet_name].iter_rows(min_row=1, max_col=len(headers)))
        header = next(rows)
        for column, expected in enumerate(headers, start=1):
            _assert_cell(header[column - 1].value, expected, f"{question}/{sheet_name} header {column}")
        for row_index, (time_s, expected_values) in enumerate(zip(times, payload[key]), start=2):
            row = next(rows)
            _assert_cell(row[0].value, time_s, f"{question}/{sheet_name} row {row_index} time")
            for column, expected in enumerate(expected_values, start=2):
                _assert_cell(row[column - 1].value, expected,
                             f"{question}/{sheet_name} row {row_index} column {column}")
        try:
            next(rows)
        except StopIteration:
            pass
        else:
            raise AssertionError(f"{question}/{sheet_name}: unexpected extra workbook rows")
    book.close()
    return {"values_match_payload": True, "rows": len(times), "columns": len(headers) - 1}


def main() -> None:
    checks = {
        "q1": _payload_check("q1", 1800, ("T", "C")),
        "q2": _payload_check("q2", 10800, ("T", "C")),
        "q3": _payload_check("q3", 3461, ("C",)),
        "q4": _payload_check("q4", 3077, ("C",)),
        "q1_workbook": _workbook_shape(RESULTS / "q1" / "result1.xlsx", ["温度", "水分浓度"], 1800, 22),
        "q2_workbook": _workbook_shape(RESULTS / "q2" / "result2.xlsx", ["温度", "水分浓度"], 10800, 22),
        "q3_workbook": _workbook_shape(RESULTS / "q3" / "result3.xlsx", ["Sheet1"], 3461, 22),
        "q4_workbook": _workbook_shape(RESULTS / "q4" / "result4.xlsx", ["Sheet1"], 3077, 22),
        "q1_workbook_values": _workbook_matches_payload(
            "q1", [("温度", "T"), ("水分浓度", "C")],
            ["时间\\到药材中心的距离", *[round(i / 10, 1) for i in range(21)]],
        ),
        "q2_workbook_values": _workbook_matches_payload(
            "q2", [("温度", "T"), ("水分浓度", "C")],
            ["时间\\到药材中心的距离", *[round(i / 10, 1) for i in range(21)]],
        ),
        "q3_workbook_values": _workbook_matches_payload(
            "q3", [("Sheet1", "C")],
            ["时间\\到药材中心的距离", *[round(i / 10, 1) for i in range(21)]],
        ),
        "q4_workbook_values": _workbook_matches_payload(
            "q4", [("Sheet1", "C")],
            ["时间\\到药材中心的距离", *[round(i / 10, 1) for i in range(20)], "药材表面"],
        ),
    }
    for question, validation_key in [("q1", "selected_boundary_method"), ("q2", "model"), ("q3", "event_h"), ("q4", "event_h")]:
        validation = json.loads((RESULTS / question / "validation.json").read_text(encoding="utf-8"))
        if validation_key not in validation:
            raise AssertionError(f"{question}: missing validation field {validation_key}")
    q4 = json.loads((RESULTS / "q4" / "validation.json").read_text(encoding="utf-8"))
    if q4.get("radius_extrapolation_used"):
        raise AssertionError("q4: radius extrapolation unexpectedly used")
    with np.load(RESULTS / "q3" / "result3_full_precision.npz") as q3_full:
        if not float(np.max(q3_full["C"][-1])) < 0.15:
            raise AssertionError("q3: final state is not strictly below the threshold")
    with np.load(RESULTS / "q4" / "result4_full_precision.npz") as q4_full:
        if not float(np.nanmax(q4_full["C"][-1])) < 0.15:
            raise AssertionError("q4: final state is not strictly below the threshold")
    checks["input_sha256"] = {
        "附件1.xlsx": _hash(DATA / "附件1.xlsx"),
        "附件2.xlsx": _hash(DATA / "附件2.xlsx"),
    }
    source_files = [PACKAGE / "run_all.py", *sorted((PACKAGE / "code").rglob("*.py"))]
    checks["source_sha256"] = {
        path.relative_to(PACKAGE).as_posix(): _hash(path) for path in source_files
    }

    q2_sensitivity = json.loads(
        (RESULTS / "q2" / "parameter_sensitivity.json").read_text(encoding="utf-8"))
    for relative, expected in q2_sensitivity["source_sha256"].items():
        if checks["source_sha256"].get(relative) != expected:
            raise AssertionError(f"q2 sensitivity was not generated by current {relative}")

    generated_files = []
    for question, number in (("q1", "1"), ("q2", "2"), ("q3", "3"), ("q4", "4")):
        folder = RESULTS / question
        generated_files.extend([
            folder / f"result{number}.xlsx",
            folder / f"result{number}_data.json",
            folder / f"result{number}_full_precision.npz",
            folder / "validation.json",
            folder / "结果表.md",
        ])
    generated_files.extend([
        RESULTS / "q1" / "边界拟合与敏感性分析.md",
        RESULTS / "q2" / "parameter_sensitivity.json",
        RESULTS / "q2" / "validation_export.json",
        RESULTS / "q3" / "validation_export.json",
        RESULTS / "q4" / "validation_export.json",
        RESULTS / "q4" / "radius_method_comparison.csv",
    ])
    generated_files.extend(sorted(RESULTS.glob("q?/*.png")))
    missing = [str(path) for path in generated_files if not path.exists()]
    if missing:
        raise AssertionError(f"Missing generated outputs: {missing}")
    checks["generated_output_sha256"] = {
        path.relative_to(PACKAGE).as_posix(): _hash(path) for path in generated_files
    }

    copied_figures = {}
    for result_figure in sorted(RESULTS.glob("q?/*.png")):
        destination = PACKAGE / "figures" / result_figure.parent.name / result_figure.name
        if not destination.exists() or _hash(destination) != _hash(result_figure):
            raise AssertionError(f"Paper figure is not synchronized: {destination}")
        copied_figures[destination.relative_to(PACKAGE).as_posix()] = _hash(destination)
    checks["synchronized_paper_figures_sha256"] = copied_figures

    paper_text = (PACKAGE / "paper.md").read_text(encoding="utf-8")
    local_images = sorted(set(re.findall(r"!\[[^\]]*\]\(([^)]+)\)", paper_text)))
    missing_images = [reference for reference in local_images
                      if not (PACKAGE / reference).exists()]
    if missing_images:
        raise AssertionError(f"Paper references missing images: {missing_images}")
    checks["paper_image_references"] = {
        "count": len(local_images), "all_exist": True, "paths": local_images
    }

    # These files are upstream data-preprocessing evidence or manually authored
    # geometry diagrams. They do not depend on the PDE face-flux implementation,
    # so the formal run validates and records them without pretending to rebuild them.
    boundary_cv_path = RESULTS / "q1" / "boundary_cv.csv"
    endpoint_path = RESULTS / "q1" / "boundary_endpoint_comparison.csv"
    boundary_cv = _csv_records(boundary_cv_path)
    endpoints = _csv_records(endpoint_path)
    expected_methods = {"linear", "pchip", "akima", "cubic_spline",
                        "smooth_spline", "stretched_exp"}
    if len(boundary_cv) != 12 or {row["series"] for row in boundary_cv} != {"T", "C"}:
        raise AssertionError("Unexpected Q1 boundary_cv.csv structure")
    if {row["method"] for row in boundary_cv} != expected_methods:
        raise AssertionError("Q1 boundary_cv.csv does not contain all six methods")
    if len(endpoints) != 84 or {row["series"] for row in endpoints} != {"T", "C"}:
        raise AssertionError("Unexpected Q1 endpoint-comparison structure")
    if {row["method"] for row in endpoints} != expected_methods:
        raise AssertionError("Q1 endpoint comparison does not contain all six methods")
    checks["upstream_preprocessing_diagnostics"] = {
        "boundary_cv_rows": len(boundary_cv),
        "endpoint_comparison_rows": len(endpoints),
        "schema_complete": True,
        "input_dependency": "附件1.xlsx (unchanged hash recorded above)",
    }

    static_assets = [
        boundary_cv_path,
        endpoint_path,
        PACKAGE / "figures" / "图03_圆柱体与dr示意图.pdf",
        PACKAGE / "figures" / "图03_圆柱体与dr示意图.png",
        PACKAGE / "figures" / "图04_有限体积离散示意图.pdf",
        PACKAGE / "figures" / "图04_有限体积离散示意图.png",
    ]
    checks["static_assets_sha256"] = {
        path.relative_to(PACKAGE).as_posix(): _hash(path) for path in static_assets
    }
    output = RESULTS / "reproducibility_check.json"
    output.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(checks, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
