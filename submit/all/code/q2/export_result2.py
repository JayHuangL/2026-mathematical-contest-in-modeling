"""Export the computed Question 2 arrays into the supplied workbook template."""
from pathlib import Path
import json
from copy import copy

import openpyxl


HERE = Path(__file__).resolve().parent
PACKAGE = HERE.parent.parent
DATA = PACKAGE / "results" / "q2" / "result2_data.json"
TEMPLATE = PACKAGE / "data" / "result2_template.xlsx"
OUTPUT = PACKAGE / "results" / "q2" / "result2.xlsx"


def main():
    payload = json.loads(DATA.read_text(encoding="utf-8"))
    times = payload["t_s"]
    distances = payload["r_cm"]
    if len(distances) != 21 or len(times) != 10800:
        raise ValueError("Question 2 requires 21 positions and 10800 one-second rows.")

    workbook = openpyxl.load_workbook(TEMPLATE)
    if workbook.sheetnames != ["温度", "水分浓度"]:
        raise ValueError(f"Unexpected template sheets: {workbook.sheetnames}")

    for sheet_name, key in (("温度", "T"), ("水分浓度", "C")):
        sheet = workbook[sheet_name]
        sheet.cell(1, 1).value = "时间\\到药材中心的距离"
        for column, distance in enumerate(distances, start=2):
            cell = sheet.cell(1, column)
            cell.value = distance
            cell.number_format = "0.0"
        for row_number, (time_s, values) in enumerate(zip(times, payload[key]), start=2):
            sheet.cell(row_number, 1).value = time_s
            sheet.cell(row_number, 1).number_format = "0"
            if len(values) != 21:
                raise ValueError(f"{sheet_name} row {row_number} does not have 21 values.")
            for column, value in enumerate(values, start=2):
                cell = sheet.cell(row_number, column)
                cell.value = value
                cell.number_format = "0.0000"
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 24
        for column in range(2, 23):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 10
        for cell in sheet[1][:23]:
            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            alignment.vertical = "center"
            alignment.wrap_text = True
            cell.alignment = alignment

    workbook.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    main()
