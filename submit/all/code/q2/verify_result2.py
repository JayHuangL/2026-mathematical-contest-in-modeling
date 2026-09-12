"""Read-only verification of the Q2 workbook against full-precision output."""
from pathlib import Path
import hashlib
import json

import numpy as np
import openpyxl


PACKAGE = Path(__file__).resolve().parents[2]
OUT = PACKAGE / "results" / "q2"
summary = json.loads((OUT / "validation.json").read_text(encoding="utf-8"))
input_path = PACKAGE / "data" / "附件1.xlsx"
assert hashlib.sha256(input_path.read_bytes()).hexdigest() == summary["input_sha256"]
data = json.loads((OUT / "result2_data.json").read_text(encoding="utf-8"))
book = openpyxl.load_workbook(OUT / "result2.xlsx", read_only=True, data_only=True)
assert book.sheetnames == ["温度", "水分浓度"]
checks = []
with np.load(OUT / "result2_full_precision.npz") as full:
    for name, key in [("温度", "T"), ("水分浓度", "C")]:
        expected = np.round(full[key][1:], 4)
        assert np.array_equal(expected, np.asarray(data[key]))
        rows = iter(book[name].iter_rows())
        header = next(rows)
        assert len(header) == 22
        assert [cell.value for cell in header[1:]] == data["r_cm"]
        count = 0
        for index, row in enumerate(rows):
            assert row[0].value == index + 1 == data["t_s"][index]
            assert len(row) == 22
            assert np.array_equal(np.asarray([cell.value for cell in row[1:]]), expected[index])
            assert all(cell.number_format == "0.0000" for cell in row[1:])
            count += 1
        assert count == 10800
        checks.append({"sheet": name, "time_rows": count, "distance_columns": 21,
                       "numeric_values_verified": count * 21})
book.close()

tables = (OUT / "结果表.md").read_text(encoding="utf-8")
for key in ("T", "C"):
    expected_table = np.round(np.asarray(summary[f"table_{key}"]), 4)
    with np.load(OUT / "result2_full_precision.npz") as full:
        assert np.array_equal(expected_table, np.round(full[key][1800::1800, ::5], 4))
    for time_h, row in zip(summary["table_times_h"], expected_table):
        line = "| " + f"{time_h:.1f}" + " | " + " | ".join(f"{value:.4f}" for value in row) + " |"
        assert line in tables

result = {"checks": checks, "all_453600_values_match": True,
          "markdown_tables_match": True, "input_unchanged": True}
(OUT / "validation_export.json").write_text(
    json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False), flush=True)
