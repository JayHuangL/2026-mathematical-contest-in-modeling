"""Export the Question 1 JSON payload to the required result1.xlsx workbook.

Run this file through submit/all/run_all.py or directly from the integrated
package. The template is kept in data/result1_template.xlsx.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


HERE = Path(__file__).resolve().parent
PACKAGE = HERE.parent.parent
if str(HERE.parent) not in sys.path:
    sys.path.insert(0, str(HERE.parent))
from artifact_names import artifact_name
TEMPLATE = PACKAGE / "data" / "result1_template.xlsx"
PAYLOAD = PACKAGE / "results" / "q1" / artifact_name("q1", "result_data")
OUTPUT = PACKAGE / "results" / "q1" / artifact_name("q1", "workbook")


def write_sheet(ws, times, positions, values):
    """Write one field in the problem's time-by-radius layout."""
    if len(times) != len(values):
        raise ValueError(f"{ws.title}: time and value row counts differ")
    if not values or len(values[0]) != len(positions):
        raise ValueError(f"{ws.title}: radius and value column counts differ")

    ws.cell(1, 1, "时间\\到药材中心的距离")
    for col, radius_cm in enumerate(positions, start=2):
        cell = ws.cell(1, col, float(radius_cm))
        cell.number_format = "0.0"

    for row, (time_s, values_at_time) in enumerate(zip(times, values), start=2):
        time_cell = ws.cell(row, 1, int(time_s))
        time_cell.number_format = "0"
        for col, value in enumerate(values_at_time, start=2):
            value_cell = ws.cell(row, col, float(value))
            value_cell.number_format = "0.0000"

    ws.freeze_panes = "B2"


def main():
    with PAYLOAD.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)

    times = payload["t_s"]
    positions = payload["r_cm"]
    temperature = payload["T"]
    moisture = payload["C"]
    if len(times) != 1800 or len(positions) != 21:
        raise ValueError("The payload must contain 1800 times and 21 radial positions.")

    workbook = openpyxl.load_workbook(TEMPLATE)
    if workbook.sheetnames[:2] != ["温度", "水分浓度"]:
        raise ValueError("The template must contain 温度 and 水分浓度 sheets first.")
    write_sheet(workbook["温度"], times, positions, temperature)
    write_sheet(workbook["水分浓度"], times, positions, moisture)
    workbook.save(OUTPUT)
    workbook.close()
    print(f"wrote {OUTPUT} ({len(times)} rows x {len(positions)} radial positions)")


if __name__ == "__main__":
    main()
