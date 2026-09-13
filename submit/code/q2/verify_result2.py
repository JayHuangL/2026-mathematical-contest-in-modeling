"""Read-only verification of the Q2 workbook against full-precision output."""
from pathlib import Path
import hashlib
import json
import sys

import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
PACKAGE = HERE.parents[1]
if str(HERE.parent) not in sys.path:
    sys.path.insert(0, str(HERE.parent))
from artifact_names import artifact_name
OUT = PACKAGE / "results" / "q2"
summary = json.loads((OUT / artifact_name("q2", "validation")).read_text(encoding="utf-8"))
input_path = PACKAGE / "data" / "附件1.xlsx"
assert hashlib.sha256(input_path.read_bytes()).hexdigest() == summary["input_sha256"]
data = json.loads((OUT / artifact_name("q2", "result_data")).read_text(encoding="utf-8"))
book = openpyxl.load_workbook(OUT / artifact_name("q2", "workbook"), read_only=True, data_only=True)
assert book.sheetnames == ["温度", "水分浓度"]
checks = []
with np.load(OUT / artifact_name("q2", "full_precision")) as full:
    for name, key in [("温度", "T"), ("水分浓度", "C")]:
        expected = np.round(full[key][1:], 4)
        assert np.array_equal(expected, np.asarray(data[key]))
        rows = iter(book[name].iter_rows())
        header = next(rows)
        assert len(header) == 22
        assert [cell.value for cell in header[1:]] == data["r_cm"]
        count = 0
        for index, row in enumerate(rows):
            assert row[0].value == index + 1 == data["t_s"][index]
            assert len(row) == 22
            assert np.array_equal(np.asarray([cell.value for cell in row[1:]]), expected[index])
            assert all(cell.number_format == "0.0000" for cell in row[1:])
            count += 1
        assert count == 10800
        checks.append({"sheet": name, "time_rows": count, "distance_columns": 21,
                       "numeric_values_verified": count * 21})
book.close()

result = {"checks": checks, "all_453600_values_match": True,
          "input_unchanged": True}
(OUT / artifact_name("q2", "validation_export")).write_text(
    json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False), flush=True)
