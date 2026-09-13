from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

import openpyxl


HERE = Path(__file__).resolve().parent
PACKAGE = HERE.parent.parent
if str(HERE.parent) not in sys.path:
    sys.path.insert(0, str(HERE.parent))
from artifact_names import artifact_name, workbook_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Build result4.xlsx from q4_result_data.json.")
    parser.add_argument("--root", type=Path, default=PACKAGE / "data")
    parser.add_argument("--output", type=Path,
                        default=workbook_path(PACKAGE / "results", "q4"))
    args = parser.parse_args()

    data = json.loads((PACKAGE / "results" / "q4" / artifact_name("q4", "result_data")).read_text(encoding="utf-8"))
    template = PACKAGE / "data" / "result4_template.xlsx"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, args.output)

    book = openpyxl.load_workbook(args.output)
    sheet = book["Sheet1"]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)

    headers = ["时间\\到药材中心的距离", *data["fixed_r_cm"], "药材表面"]
    for column, value in enumerate(headers, start=1):
        sheet.cell(1, column, value)

    for row_index, (time_s, values) in enumerate(zip(data["t_s"], data["C"]), start=2):
        sheet.cell(row_index, 1, time_s)
        for column, value in enumerate(values, start=2):
            sheet.cell(row_index, column, value)

    last_row = len(data["t_s"]) + 1
    for row in sheet.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=len(headers)):
        row[0].number_format = "0.0000" if row[0].row == last_row else "0"
        for cell in row[1:]:
            cell.number_format = "0.0000"
    for cell in sheet[1][: len(headers)]:
        cell.number_format = "0.0" if cell.column > 1 else "General"
    sheet.freeze_panes = "B2"
    book.save(args.output)
    print(f"Saved {args.output}")


if __name__ == "__main__":
    main()
